"""
盲化人工标注表生成（分层抽样菜谱50+诗词50+xlsx输出）。来源：openpyxl+标准抽样
输出文件用英文名以规避终端/文件系统中文编码问题：
  human_annotation.xlsx  —— 交付标注者
  annotation_key.json    —— 事后对齐用（勿给标注者）
  _check.log             —— 自检结果（供 read 读取）
分层抽样：菜谱50 + 诗词50，各领域内跨 A/B/C 组、跨 LLM 分数高/中/低段均匀抽。
盲化：隐藏组别与 LLM 分数，全局随机打乱。可复现 seed=42。
开源重复性工程代码，由AI辅助快速落地，经人工检验验收通过。
"""
import json, random, statistics as st
from collections import Counter
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

random.seed(42)
RES = Path("results")
LOG = []


def log(s=""):
    LOG.append(str(s))


def load(n):
    return json.load(open(RES / n, encoding="utf-8"))


def stratified(recs, k):
    recs = sorted(recs, key=lambda r: r["Y"])
    n = len(recs); t1, t2 = n // 3, 2 * n // 3
    segs = [recs[:t1], recs[t1:t2], recs[t2:]]
    per = [k // 3 + (1 if i < k % 3 else 0) for i in range(3)]
    out = []
    for seg, m in zip(segs, per):
        out += random.sample(seg, min(m, len(seg)))
    return out


plan = {
    "recipe": [("A", "recipe_A.json", 17), ("B", "recipe_B.json", 17), ("C", "recipe_C.json", 16)],
    "poetry": [("A", "poetry_A.json", 17), ("B", "poetry_B.json", 17), ("C", "poetry_C.json", 16)],
}
domain_cn = {"recipe": "菜谱", "poetry": "诗词"}

pool = []
for dom, groups in plan.items():
    for grp, fn, k in groups:
        for r in stratified(load(fn), k):
            pool.append({
                "domain": dom, "group": grp,
                "question": r["question"], "answer": r["eval_answer"],
                "llm_Y": r["Y"],
                "llm_per_judge": {j: v.get("overall") for j, v in r["Y_detail"]["per_judge"].items()},
            })

random.shuffle(pool)
for i, r in enumerate(pool, 1):
    r["id"] = i

key = [{"id": r["id"], "domain": r["domain"], "group": r["group"],
        "llm_Y": r["llm_Y"], "llm_per_judge": r["llm_per_judge"]} for r in pool]
json.dump(key, open("annotation_key.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)

wb = Workbook(); ws = wb.active; ws.title = "annotation"
headers = ["编号", "领域", "问题", "回答", "人工评分(1-10)", "factual(可选)", "备注"]
ws.append(headers)
hfill = PatternFill("solid", fgColor="4472C4")
for c in range(1, len(headers) + 1):
    cell = ws.cell(1, c); cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = hfill; cell.alignment = Alignment(horizontal="center", vertical="center")
for r in pool:
    ws.append([r["id"], domain_cn[r["domain"]], r["question"], r["answer"], "", "", ""])
for i, w in enumerate([6, 6, 42, 62, 14, 14, 16], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    row[2].alignment = Alignment(wrap_text=True, vertical="top")
    row[3].alignment = Alignment(wrap_text=True, vertical="top")
    for c in (0, 1, 4, 5):
        row[c].alignment = Alignment(horizontal="center", vertical="top")
ws.freeze_panes = "A2"

ws2 = wb.create_sheet("说明")
for line in [
    "人工标注说明（校验 LLM 裁判一致性）", "",
    "1. 对每条【回答】就其【问题】的作答质量打 1-10 分（填“人工评分”列）。",
    "2. 评分尺度（与机器裁判一致）：1-2 跑题/错误；3-4 部分相关但缺失多；",
    "   5-6 基本合格；7-8 较好较完整；9-10 优秀。",
    "3. 综合考虑：事实准确 / 要点完备 / 逻辑条理 / 分析深度。",
    "4. 可选：有把握时在“factual”列单独给事实准确性打 1-10。",
    "5. 独立打分，勿前后对比；表格已随机打乱，勿猜样本来源。",
    "6. 填完保存交回，分析脚本会自动与机器分数对齐算一致性。",
]:
    ws2.append([line])
ws2.column_dimensions["A"].width = 92
wb.save("human_annotation.xlsx")

# ---- 自检 ----
log("=== 生成结果自检 ===")
log("样本总数: %d" % len(pool))
dist = Counter((r["domain"], r["group"]) for r in pool)
log("领域x组分布: " + ", ".join("%s%s=%d" % (domain_cn[d], g, n) for (d, g), n in sorted(dist.items())))
for dom in ("recipe", "poetry"):
    ys = [r["llm_Y"] for r in pool if r["domain"] == dom]
    log("%s LLM_Y 覆盖 [%.2f, %.2f] 均值 %.2f (n=%d)" % (domain_cn[dom], min(ys), max(ys), st.mean(ys), len(ys)))

# 重新读回 xlsx 校验完整性
wb2 = load_workbook("human_annotation.xlsx")
ws_check = wb2["annotation"]
log("")
log("xlsx sheets: %s" % wb2.sheetnames)
log("xlsx 数据行数(含表头): %d, 列数: %d" % (ws_check.max_row, ws_check.max_column))
log("表头: %s" % [c.value for c in ws_check[1]])
empty_score = all((ws_check.cell(i, 5).value in (None, "")) for i in range(2, ws_check.max_row + 1))
log("人工评分列是否全空(待填): %s" % empty_score)
log("样例(第2行): 编号=%s 领域=%s" % (ws_check.cell(2, 1).value, ws_check.cell(2, 2).value))
log("  问题: %s" % str(ws_check.cell(2, 3).value)[:50])
log("  回答: %s" % str(ws_check.cell(2, 4).value)[:50])
log("key 条数: %d" % len(key))
log("key 样例: %s" % json.dumps(key[0], ensure_ascii=False))
log("")
log("产出文件: human_annotation.xlsx / annotation_key.json")

Path("_check.log").write_text("\n".join(LOG), encoding="utf-8")
