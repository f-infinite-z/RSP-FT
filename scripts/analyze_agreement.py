"""
人机评分一致性（Pearson/Spearman/ICC+分领域分组均值对照）。来源：ICC(2,1)(Shrout&Fleiss,1979)
读取：human_annotation.xlsx（第5列“人工评分”，可选第6列 factual）+ annotation_key.json
计算：人工 vs LLM 的 Pearson / Spearman / ICC(2,1) / 平均绝对偏差；
      分领域 + A/B/C 组人机均值对照（看排序方向一致性）。
产出：results/agreement_report.txt（同时写 _agreement.log 供快速查看）
用法：python analyze_agreement.py

开源重复性工程代码，由AI辅助快速落地，经人工检验验收通过。
"""
import json, math, statistics as st
from pathlib import Path
from openpyxl import load_workbook

L = []
def log(s=""): L.append(str(s))

def pearson(xs, ys):
    mx, my = st.mean(xs), st.mean(ys)
    cov = sum((x-mx)*(y-my) for x, y in zip(xs, ys))
    sx = math.sqrt(sum((x-mx)**2 for x in xs)); sy = math.sqrt(sum((y-my)**2 for y in ys))
    return cov/(sx*sy) if sx and sy else float("nan")

def spearman(xs, ys):
    def rank(v):
        order = sorted(range(len(v)), key=lambda i: v[i]); r = [0.0]*len(v)
        i = 0
        while i < len(v):
            j = i
            while j+1 < len(v) and v[order[j+1]] == v[order[i]]:
                j += 1
            avg = (i+j)/2.0
            for t in range(i, j+1):
                r[order[t]] = avg
            i = j+1
        return r
    return pearson(rank(xs), rank(ys))

def icc21(pairs):
    """ICC(2,1) 双向随机、单测量、绝对一致性。pairs=[(human, llm), ...]"""
    n = len(pairs); k = 2
    grand = sum(a+b for a, b in pairs)/(n*k)
    row_means = [(a+b)/2 for a, b in pairs]
    col_means = [sum(p[0] for p in pairs)/n, sum(p[1] for p in pairs)/n]
    MSR = k*sum((rm-grand)**2 for rm in row_means)/(n-1)
    MSC = n*sum((cm-grand)**2 for cm in col_means)/(k-1)
    SST = sum((v-grand)**2 for p in pairs for v in p)
    SSR = k*sum((rm-grand)**2 for rm in row_means)
    SSC = n*sum((cm-grand)**2 for cm in col_means)
    SSE = SST - SSR - SSC
    MSE = SSE/((n-1)*(k-1))
    denom = MSR + (k-1)*MSE + (k/n)*(MSC-MSE)
    return (MSR-MSE)/denom if denom else float("nan")

wb = load_workbook("human_annotation.xlsx")
ws = wb["annotation"]
key = {r["id"]: r for r in json.load(open("annotation_key.json", encoding="utf-8"))}

def _isnum(x):
    try:
        float(x); return True
    except (TypeError, ValueError):
        return False

EXCLUDE_KW = ["无效", "信息不足", "主观", "剔除", "重名", "词牌", "答非所问-题"]

rows = []
n_unlabeled = 0
n_excluded = 0
excluded_ids = []
for i in range(2, ws.max_row+1):
    hid = ws.cell(i, 1).value
    hs = ws.cell(i, 5).value
    fac = ws.cell(i, 6).value
    note = ws.cell(i, 7).value or ""
    if not _isnum(hs):
        n_unlabeled += 1
        continue
    # 备注含排除关键词 → 剔除（如词牌重名/主观偏好题）
    if any(kw in str(note) for kw in EXCLUDE_KW):
        n_excluded += 1
        excluded_ids.append(hid)
        continue
    k = key.get(hid)
    if not k:
        continue
    rows.append({"id": hid, "domain": k["domain"], "group": k["group"],
                 "human": float(hs), "llm": k["llm_Y"],
                 "human_fac": float(fac) if _isnum(fac) else None})

log("=== 人机评分一致性分析 ===")
log("有效标注: %d 条 | 剔除(备注无效): %d 条 | 未标注: %d 条 | 总计: %d" %
    (len(rows), n_excluded, n_unlabeled, ws.max_row - 1))
if excluded_ids:
    log("剔除的编号: %s" % excluded_ids)
if len(rows) < 5:
    log("有效标注不足（需≥5），请先在 human_annotation.xlsx 的“人工评分”列填分。")
    Path("results").mkdir(exist_ok=True)
    Path("results/agreement_report.txt").write_text("\n".join(L), encoding="utf-8")
    Path("_agreement.log").write_text("\n".join(L), encoding="utf-8")
    raise SystemExit

def block(name, rs):
    h = [r["human"] for r in rs]; m = [r["llm"] for r in rs]
    log("")
    log("[%s] n=%d" % (name, len(rs)))
    log("  Pearson r  = %.3f" % pearson(h, m))
    log("  Spearman ρ = %.3f" % spearman(h, m))
    log("  ICC(2,1)   = %.3f" % icc21(list(zip(h, m))))
    log("  平均绝对偏差 |人-机| = %.2f  (人均值 %.2f vs 机均值 %.2f)" %
        (st.mean([abs(a-b) for a, b in zip(h, m)]), st.mean(h), st.mean(m)))

block("全部", rows)
for dom, cn in [("recipe", "菜谱"), ("poetry", "诗词")]:
    rs = [r for r in rows if r["domain"] == dom]
    if len(rs) >= 5:
        block(cn, rs)

log("")
log("=== A/B/C 组人机均值对照（排序方向一致性）===")
for dom, cn in [("recipe", "菜谱"), ("poetry", "诗词")]:
    log("[%s]" % cn)
    for g in "ABC":
        rs = [r for r in rows if r["domain"] == dom and r["group"] == g]
        if rs:
            log("  %s组: 人工 %.2f | 机器 %.2f (n=%d)" %
                (g, st.mean([r["human"] for r in rs]), st.mean([r["llm"] for r in rs]), len(rs)))

fac = [(r["human_fac"], key[r["id"]]["llm_per_judge"]) for r in rows if r.get("human_fac") is not None]
if len(fac) >= 5:
    hf = [f[0] for f in fac]
    lf = [st.mean([v for v in f[1].values() if v is not None]) for f in fac]
    log("")
    log("[factual 可选] n=%d  Pearson r=%.3f  Spearman ρ=%.3f" % (len(fac), pearson(hf, lf), spearman(hf, lf)))

log("")
log("判读：Pearson/Spearman ≥0.6 或 ICC ≥0.6 即可支持“LLM 裁判与人工判断一致”，")
log("      结合 A/B/C 排序方向一致，可压制“评测偏见/裁判不可信”质疑。")

Path("results").mkdir(exist_ok=True)
Path("results/agreement_report.txt").write_text("\n".join(L), encoding="utf-8")
Path("_agreement.log").write_text("\n".join(L), encoding="utf-8")
